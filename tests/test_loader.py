import importlib.util
import os
import unittest

from rag.ingest.loader import Document, iter_files, load_documents, load_one

try:  # 兼容 python -m unittest discover -s tests 与 -t . 两种调用方式
    from tests.support import PROJECT_ROOT, TempWorkspace
except ImportError:  # pragma: no cover
    from support import PROJECT_ROOT, TempWorkspace


class DocIdTest(unittest.TestCase):
    """doc_id 必须唯一：撞车会导致 chunk_id 撞车，检索静默指向错误文档。"""

    def test_doc_ids_unique_with_conflicting_names(self):
        with TempWorkspace(prefix="loader-") as tmp:
            os.makedirs(os.path.join(tmp, "sub"), exist_ok=True)
            for relative in ("a.md", "a-1.md", os.path.join("sub", "a.md")):
                path = os.path.join(tmp, relative)
                with open(path, "w", encoding="utf-8") as handle:
                    handle.write(f"# 标题\n这是 {relative} 的内容。\n")
            docs = load_documents([tmp])
            ids = [doc.doc_id for doc in docs]
            self.assertEqual(len(ids), 3)
            self.assertEqual(len(ids), len(set(ids)), f"doc_id 撞车：{ids}")
            chunk_ids = [f"{doc_id}#0" for doc_id in ids]
            self.assertEqual(len(chunk_ids), len(set(chunk_ids)))


class LoaderTest(unittest.TestCase):
    def test_iter_files_supports_dirs_and_files(self):
        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "x.md")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("# T\n正文\n")
            self.assertEqual(iter_files([tmp]), [path])
            self.assertEqual(iter_files([path]), [path])

    def test_missing_file_raises(self):
        missing = os.path.join(PROJECT_ROOT, ".testtmp", "definitely-not-here.docx")
        with self.assertRaises(Exception):
            load_one(missing)

    def test_document_title_falls_back_to_doc_id(self):
        doc = Document(doc_id="d1", text="没有标题的正文")
        self.assertEqual(doc.title, "d1")

    def test_gb18030_text_file(self):
        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "gbk.txt")
            with open(path, "wb") as handle:
                handle.write("# 标题\n这是 GBK 编码的正文。\n".encode("gb18030"))
            text = load_one(path)
            self.assertIn("GBK 编码的正文", text)


class FormatTest(unittest.TestCase):
    """各语料格式的解析路径（pdf 需要 pypdf，此处不覆盖）。"""

    def test_jsonl_records(self):
        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "rows.jsonl")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write('{"a": 1, "b": "二"}\n{"a": 2}\n')
            text = load_one(path)
            self.assertIn("a: 1", text)
            self.assertIn("b: 二", text)

    def test_json_array_of_objects(self):
        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "rows.json")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write('[{"name": "甲", "score": 1}, {"name": "乙", "score": 2}]')
            text = load_one(path)
            self.assertIn("name: 甲", text)
            self.assertIn("score: 2", text)

    def test_csv_rows_become_key_value_pairs(self):
        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "table.csv")
            with open(path, "w", encoding="utf-8", newline="") as handle:
                handle.write("method,params,note\nBM25,k1=1.5;b=0.75,词频饱和\nRRF,k=60,只用名次\n")
            text = load_one(path)
            self.assertIn("method: BM25", text)
            self.assertIn("params: k=60", text)
            self.assertIn("note: 只用名次", text)

    @unittest.skipUnless(importlib.util.find_spec("docx") is not None, "需要 python-docx")
    def test_docx_headings_paragraphs_and_tables(self):
        import docx

        with TempWorkspace(prefix="loader-") as tmp:
            path = os.path.join(tmp, "doc.docx")
            document = docx.Document()
            document.add_heading("检索算法", level=1)
            document.add_paragraph("BM25 使用 k1 与 b 两个参数。")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "指标"
            table.rows[0].cells[1].text = "含义"
            document.save(path)

            text = load_one(path)
            self.assertIn("# 检索算法", text)  # 标题层级被还原成 Markdown
            self.assertIn("BM25 使用 k1 与 b 两个参数。", text)
            self.assertIn("指标 | 含义", text)

    @unittest.skipUnless(importlib.util.find_spec("docx") is not None, "需要 python-docx")
    def test_docx_end_to_end_through_pipeline(self):
        import docx

        from rag.config import Config
        from rag.pipeline import RAGPipeline

        with TempWorkspace(prefix="docx-pipe-") as root:
            corpus_dir = os.path.join(root, "corpus")
            os.makedirs(corpus_dir, exist_ok=True)
            document = docx.Document()
            document.add_heading("融合", level=2)
            document.add_paragraph("RRF 的常数 k 通常取 60，只使用排名。")
            document.save(os.path.join(corpus_dir, "notes.docx"))

            cfg = Config.from_dict(
                {
                    "paths": {"corpus": "corpus", "index": "index", "reports": "reports", "web": "web"},
                    "embedder": {"provider": "hashing", "dim": 128},
                    "generation": {"provider": "extractive", "extractive_min_coverage": 0.1, "extractive_min_hits": 2},
                    "logging": {"level": "ERROR"},
                },
                root=root,
            )
            pipeline = RAGPipeline(cfg)
            report = pipeline.ingest()
            self.assertEqual(report.docs, 1)
            answer = pipeline.ask("RRF 的常数 k 取多少？")
            self.assertIn("notes", answer.debug["ranked_docs"][0])


if __name__ == "__main__":
    unittest.main()
