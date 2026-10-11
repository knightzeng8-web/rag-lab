"""重依赖自检：装完 venv 后一条命令确认 BGE / Cross-Encoder / faiss / FastAPI 是否真的能用。

    .\run.ps1 scripts/check_heavy_deps.py            # 全部检查（首次会下载模型）
    .\run.ps1 scripts/check_heavy_deps.py --quick    # 只查 import 与版本，不下载模型

每项都会打印 PASS/FAIL 与耗时；任何一项 FAIL 都以非零退出码结束，可以直接当 CI 关卡用。
"""

from __future__ import annotations

import argparse
import os
import sys
import time
import traceback
from typing import Callable, List, Tuple

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)


def _timed(fn: Callable[[], str]) -> Tuple[bool, str, float]:
    started = time.perf_counter()
    try:
        detail = fn()
        return True, detail or "", (time.perf_counter() - started) * 1000.0
    except Exception as exc:  # noqa: BLE001 - 自检脚本要兜住一切
        return False, f"{type(exc).__name__}: {exc}", (time.perf_counter() - started) * 1000.0


def check_pyyaml() -> str:
    from rag.config import load_config

    cfg = load_config()
    assert cfg.path is not None and cfg.path.endswith((".yaml", ".yml")), f"未从 YAML 载入：{cfg.path}"
    return f"config={os.path.basename(cfg.path)}  embedder.device={cfg.get('embedder.device')}"


def check_torch() -> str:
    import torch

    cuda = torch.cuda.is_available()
    gpu = torch.cuda.get_device_name(0) if cuda else "-"
    return f"torch={torch.__version__}  cuda={cuda}  gpu={gpu}"


def check_bge() -> str:
    from rag.index.embedder import BGEEmbedder

    embedder = BGEEmbedder()
    vectors = embedder.encode(["混合检索使用 RRF 融合", "BM25 使用 k1 与 b"])
    return (
        f"model={embedder.model_name}  device={embedder.device}  dim={vectors.shape[1]}  "
        f"shape={vectors.shape}"
    )


def check_cross_encoder() -> str:
    from rag.retrieve.rerank import CrossEncoderReranker
    from rag.types import Chunk, Scored

    reranker = CrossEncoderReranker()
    candidates = [
        Scored(Chunk(chunk_id="d#0", doc_id="d", text="完全无关的内容：今天天气不错。"), 0.9),
        Scored(Chunk(chunk_id="d#1", doc_id="d", text="RRF 使用倒数排名融合，常数 k 取 60。"), 0.1),
    ]
    ranked = reranker.rerank("RRF 的常数 k 取多少？", candidates, top_k=2)
    assert ranked[0].chunk_id == "d#1", f"重排未把相关片段排到第一：{[r.chunk_id for r in ranked]}"
    return f"model={reranker.model_name}  device={reranker.device}  top1={ranked[0].chunk_id}"


def check_faiss() -> str:
    import faiss
    import numpy as np

    from rag.index.store import FaissVectorStore
    from rag.types import Chunk

    store = FaissVectorStore(dim=8)
    chunks = [Chunk(chunk_id=f"d#{i}", doc_id="d", text=str(i)) for i in range(4)]
    matrix = np.eye(4, 8, dtype=np.float32)
    store.add(chunks, matrix)
    hits = store.search(matrix[2], top_n=1)
    assert hits and hits[0].chunk_id == "d#2", f"faiss 检索结果不对：{hits}"
    return f"faiss={getattr(faiss, '__version__', '?')}  backend={store.backend}  top1={hits[0].chunk_id}"


def check_fastapi() -> str:
    import fastapi
    import uvicorn

    from rag.api.main import create_app  # noqa: F401 - 只验证可导入

    return f"fastapi={fastapi.__version__}  uvicorn={uvicorn.__version__}"


def check_docx() -> str:
    import docx

    from rag.ingest.loader import load_one
    from tests.support import TempWorkspace

    with TempWorkspace(prefix="heavy-") as tmp:
        path = os.path.join(tmp, "x.docx")
        document = docx.Document()
        document.add_heading("标题", level=1)
        document.add_paragraph("正文。")
        document.save(path)
        text = load_one(path)
    assert "标题" in text and "正文。" in text
    return f"python-docx={docx.__version__ if hasattr(docx, '__version__') else '?'}"


def check_pypdf() -> str:
    import pypdf

    return f"pypdf={pypdf.__version__}"


CHECKS: List[Tuple[str, Callable[[], str], bool]] = [
    ("pyyaml（读 YAML 配置）", check_pyyaml, False),
    ("torch（含 CUDA 检测）", check_torch, False),
    ("BGE 向量模型", check_bge, True),
    ("Cross-Encoder 重排", check_cross_encoder, True),
    ("faiss 向量检索", check_faiss, False),
    ("FastAPI + uvicorn", check_fastapi, False),
    ("python-docx 解析", check_docx, False),
    ("pypdf 解析", check_pypdf, False),
]


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="rag-lab 重依赖自检")
    parser.add_argument("--quick", action="store_true", help="跳过需要下载模型的两项")
    args = parser.parse_args(argv)

    print(f"Python {sys.version.split()[0]}；HF_HOME={os.environ.get('HF_HOME', '(默认)')}")
    print(f"HF_ENDPOINT={os.environ.get('HF_ENDPOINT', '(默认)')}\n")

    results: List[Tuple[str, bool, str, float]] = []
    for name, fn, is_model in CHECKS:
        if args.quick and is_model:
            print(f"[skip] {name}（--quick）")
            continue
        print(f"[ .. ] {name} …", flush=True)
        ok, detail, ms = _timed(fn)
        results.append((name, ok, detail, ms))
        print(f"[{' OK ' if ok else 'FAIL'}] {name}  ({ms:.0f} ms)  {detail}", flush=True)
        if not ok and "Traceback" not in detail:
            continue

    print("\n================ 汇总 ================")
    failed = 0
    for name, ok, detail, ms in results:
        if not ok:
            failed += 1
        print(f"{'PASS' if ok else 'FAIL'}  {name:26s} {ms:8.0f} ms  {detail[:80]}")
    print("=====================================")
    if failed:
        print(f"{failed} 项失败。缺依赖不会导致主程序崩溃（会自动降级），但对应能力不可用。")
        return 1
    print("全部通过。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
