"""语料加载：把不同格式的文件统一成 (doc_id, text, meta)。"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from ..text import first_heading, normalize_ws, stable_hash

__all__ = ["Document", "SUPPORTED_SUFFIXES", "load_documents", "iter_files", "load_one"]

SUPPORTED_SUFFIXES = (".md", ".markdown", ".txt", ".text", ".json", ".jsonl", ".csv", ".docx", ".pdf")


@dataclass
class Document:
    doc_id: str
    text: str
    path: str = ""
    meta: Dict[str, Any] = field(default_factory=dict)

    @property
    def title(self) -> str:
        return str(self.meta.get("title") or first_heading(self.text) or self.doc_id)


def iter_files(inputs: Iterable[str]) -> List[str]:
    """入参可以是文件或目录，递归收集受支持的文件。"""
    files: List[str] = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            for suffix in SUPPORTED_SUFFIXES:
                files.extend(str(f) for f in sorted(p.rglob(f"*{suffix}")))
        elif p.is_file():
            files.append(str(p))
    # 去重且保持顺序
    seen = set()
    unique: List[str] = []
    for f in files:
        key = os.path.abspath(f)
        if key not in seen:
            seen.add(key)
            unique.append(f)
    return unique


def _doc_id_for(path: str, used: Dict[str, int]) -> str:
    """生成唯一 doc_id。

    注意不能只判断 ``stem in used``：语料里可能本来就存在 ``a.md`` 与 ``a-1.md``，
    如果生成的 ``a-1`` 不登记进 used，两者会撞成同一个 doc_id，进而产生相同的
    chunk_id —— 检索层会静默把结果指向错误的文档。所以这里对「生成出来的 id」
    也做占用检查，并把计数写回原 stem。
    """
    stem = Path(path).stem or "doc"
    counter = used.get(stem, 0)
    base = stem
    while base in used:
        counter += 1
        base = f"{stem}-{counter}"
    used[stem] = counter
    used[base] = 0
    return base


def _read_text_file(path: str) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return Path(path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def _read_docx(path: str) -> str:
    try:
        import docx  # type: ignore  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(f"读取 docx 需要 python-docx：pip install python-docx（{path}）") from exc
    document = docx.Document(path)
    blocks: List[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            blocks.append(f"{'#' * min(int(level or 1), 6)} {text}")
        else:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _read_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(f"读取 pdf 需要 pypdf：pip install pypdf（{path}）") from exc
    reader = PdfReader(path)
    pages: List[str] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _read_json_ish(path: str) -> str:
    import json

    text = _read_text_file(path)
    if path.endswith(".jsonl"):
        records = []
        for line in text.splitlines():
            line = line.strip()
            if line:
                records.append(json.loads(line))
    else:
        records = json.loads(text)
        if isinstance(records, dict):
            records = [records]
    blocks: List[str] = []
    for record in records:
        if isinstance(record, dict):
            blocks.append(
                "\n".join(f"{k}: {v}" for k, v in record.items() if str(v).strip())
            )
        else:
            blocks.append(str(record))
    return "\n\n".join(b for b in blocks if b.strip())


def _read_csv(path: str) -> str:
    import csv

    rows: List[str] = []
    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle)
        header: Optional[List[str]] = None
        for row in reader:
            if header is None:
                header = [c.strip() for c in row]
                continue
            rows.append(" | ".join(f"{h}: {c.strip()}" for h, c in zip(header, row) if c.strip()))
    return "\n".join(rows)


def load_one(path: str) -> str:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in (".json", ".jsonl"):
        return _read_json_ish(path)
    if suffix == ".csv":
        return _read_csv(path)
    return _read_text_file(path)


def load_documents(inputs: Iterable[str], log=None) -> List[Document]:
    """加载并清洗，返回 Document 列表。"""
    files = iter_files(inputs)
    used: Dict[str, int] = {}
    documents: List[Document] = []
    for path in files:
        try:
            text = normalize_ws(load_one(path))
        except Exception as exc:  # 单个文件失败不影响整体
            if log:
                log.warning("跳过无法解析的文件 %s：%s", path, exc)
            continue
        if not text:
            if log:
                log.warning("跳过空文件：%s", path)
            continue
        doc_id = _doc_id_for(path, used)
        documents.append(
            Document(
                doc_id=doc_id,
                text=text,
                path=os.path.abspath(path),
                meta={
                    "title": first_heading(text) or Path(path).stem,
                    "source": os.path.basename(path),
                    "chars": len(text),
                    "fingerprint": f"{stable_hash(text, 10**9):09d}",
                },
            )
        )
    return documents
